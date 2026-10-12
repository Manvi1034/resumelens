import io

import pytest
from docx import Document
from pypdf.errors import DependencyError

import resume_parser
from make_sample_resumes import (
    PROFILES,
    build_docx,
    build_docx_nested_table,
    build_docx_textbox,
    build_pdf,
    build_scanned_pdf,
)
from resume_parser import MAX_FILE_SIZE_MB, ResumeParseError, extract_resume_text
from skill_extractor import build_pattern, extract_skills
from text_cleaner import clean_text

GOOD_DOCX = build_docx(PROFILES["data_analyst"])
GOOD_PDF = build_pdf(PROFILES["data_analyst"])


# ---------- DOCX layouts (bug: these used to be skipped) ----------


def test_docx_skills_in_a_nested_table_are_read():
    text = extract_resume_text("resume.docx", build_docx_nested_table())
    for skill in ("Python", "Docker", "Git", "FastAPI"):
        assert skill in text


def test_docx_text_box_content_is_read():
    text = extract_resume_text("resume.docx", build_docx_textbox())
    assert "Python, SQL, Pandas, Power BI, Excel" in text


def test_docx_header_and_footer_text_is_read():
    document = Document()
    document.add_paragraph("Body text about projects and study. " * 3)
    document.sections[0].header.paragraphs[0].text = "Header skill: Kubernetes"
    document.sections[0].footer.paragraphs[0].text = "Footer skill: Tableau"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_resume_text("resume.docx", buffer.getvalue())
    assert "Kubernetes" in text and "Tableau" in text


def test_normal_docx_text_is_not_duplicated():
    text = extract_resume_text("resume.docx", GOOD_DOCX)
    assert text.count("Tableau") == 1


# ---------- broken, unusual and borderline files ----------


def test_scanned_pdf_gets_the_scanned_message():
    with pytest.raises(ResumeParseError, match="scanned"):
        extract_resume_text("scan.pdf", build_scanned_pdf())


def test_pdf_encryption_that_cannot_be_read_gets_a_clear_message(monkeypatch):
    def raise_dependency_error(*args, **kwargs):
        raise DependencyError("cryptography>=3.1 is required for AES algorithm")

    monkeypatch.setattr(resume_parser, "PdfReader", raise_dependency_error)
    with pytest.raises(ResumeParseError, match="encryption"):
        extract_resume_text("resume.pdf", b"%PDF-1.4 anything")


@pytest.mark.parametrize(
    "filename, data, message",
    [
        ("resume.pdf", b"", "empty"),
        ("resume.docx", b"", "empty"),
        ("resume.doc", b"x" * 100, "Unsupported"),
        ("resume.txt", b"x" * 100, "Unsupported"),
        ("resume", b"x" * 100, "Unsupported"),
        ("resume.pdf", GOOD_PDF[: len(GOOD_PDF) // 2], "could not be read"),
        ("resume.docx", GOOD_DOCX[: len(GOOD_DOCX) // 2], "could not be read"),
        ("resume.docx", b"PK\x03\x04" + bytes(range(256)) * 20, "could not be read"),
    ],
    ids=[
        "empty-pdf", "empty-docx", "old-doc", "txt", "no-extension",
        "truncated-pdf", "truncated-docx", "random-bytes-docx",
    ],
)
def test_bad_files_give_a_readable_message(filename, data, message):
    with pytest.raises(ResumeParseError, match=message):
        extract_resume_text(filename, data)


def test_file_names_with_capitals_spaces_and_unicode_are_accepted():
    assert extract_resume_text("My R\u00e9sum\u00e9 (final) \u65e5\u672c.DOCX", GOOD_DOCX)
    assert extract_resume_text("RESUME.PDF", GOOD_PDF)


def test_size_limit_boundary_is_exact():
    limit = MAX_FILE_SIZE_MB * 1024 * 1024
    assert extract_resume_text("big.pdf", GOOD_PDF + b"0" * (limit - len(GOOD_PDF)))
    with pytest.raises(ResumeParseError, match="larger"):
        extract_resume_text("big.pdf", GOOD_PDF + b"0" * (limit - len(GOOD_PDF) + 1))


# ---------- hyphenated skill names (bug: these used to be missed) ----------


@pytest.mark.parametrize(
    "text, skill",
    [
        ("built a machine-learning pipeline", "Machine Learning"),
        ("deep-learning models", "Deep Learning"),
        ("data-visualization dashboards", "Data Visualization"),
        ("image-processing steps", "Image Processing"),
        ("prompt-engineering experiments", "Prompt Engineering"),
        ("scikit learn and scikit-learn", "scikit-learn"),
        ("Power-BI reports", "Power BI"),
    ],
)
def test_hyphen_and_space_are_interchangeable_inside_skill_names(text, skill):
    assert skill in extract_skills(clean_text(text))


def test_joined_or_extended_words_are_still_rejected():
    skills = extract_skills(clean_text("machinelearning and machine learnings"))
    assert "Machine Learning" not in skills


def test_blank_terms_never_match_everything():
    pattern = build_pattern(["", "   ", "python"])
    assert pattern.search("python")
    assert pattern.search("java") is None