"""Create synthetic sample resumes (DOCX and PDF) for testing.

All people and details are invented. Run this file to write the resumes to the
sample_resumes folder:

    python make_sample_resumes.py

The expected answers for each resume (top role and true skills) are stored in
tests/test_cases.csv.
"""

import io
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import parse_xml
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUTPUT_FOLDER = Path(__file__).resolve().parent / "sample_resumes"

CONTACT = "sample.candidate@example.com | +91 98765 43210"
EDUCATION = (
    "B.Tech in Computer Science and Engineering, Sample Institute of Technology, "
    "2022 - 2026"
)

# Markup: "# " = name, "## " = section heading, "- " = bullet, anything else = text.
PROFILES = {
    "data_analyst": f"""
# Sample Candidate A
{CONTACT}
## Summary
Final year Computer Science student who enjoys turning raw data into clear reports and dashboards.
## Education
{EDUCATION}
## Skills
Languages: Python, SQL
Tools: Excel, Power BI, Tableau, Jupyter
Libraries: Pandas, NumPy, Matplotlib
Concepts: Statistics, data visualization
## Projects
Sales Dashboard
- Cleaned a large sales dataset with Pandas and wrote SQL queries in MySQL to summarise revenue by region
- Built an interactive Power BI dashboard used by a college club to track monthly performance
Student Survey Analysis
- Used statistics such as correlation and hypothesis testing to analyse survey results
- Presented the findings with Matplotlib charts and an Excel summary sheet
""",
    "ml_engineer": f"""
# Sample Candidate B
{CONTACT}
## Summary
Student interested in building machine learning models and deploying them as simple services.
## Education
{EDUCATION}
## Skills
Languages: Python
Machine learning: scikit-learn, Pandas, NumPy, MLflow
Deployment: FastAPI, Docker, Git, AWS
## Projects
House Price Predictor
- Trained regression models with scikit-learn and tracked the experiments with MLflow
- Wrapped the best model in a FastAPI service and packaged it with Docker
Customer Churn Model
- Compared several machine learning classifiers using cross validation and Pandas for feature engineering
- Deployed the service on AWS and kept the code in Git with a CI/CD pipeline using GitHub Actions
""",
    "nlp_engineer": f"""
# Sample Candidate C
{CONTACT}
## Summary
Student focused on language technology, from text cleaning to fine tuning transformer models.
## Education
{EDUCATION}
## Skills
Languages: Python
NLP: NLTK, spaCy, Transformers, Hugging Face
Deep learning: PyTorch
Tools: Git
## Projects
Review Sentiment Classifier
- Fine tuned a transformer model from Hugging Face on product reviews using PyTorch
- Prepared the text with NLTK and spaCy and tracked the work in Git
Question Answering Bot
- Built a chatbot that sends questions to a large language model and cleans the answers with NLP techniques
""",
    "cv_engineer": f"""
# Sample Candidate D
{CONTACT}
## Summary
Student who builds systems that understand images and video.
## Education
{EDUCATION}
## Skills
Languages: Python
Vision: OpenCV, YOLO, CNN, image processing
Libraries: PyTorch, NumPy, Matplotlib
Tools: Linux
## Projects
Helmet Detection
- Trained a YOLO model to detect helmets in traffic camera video and measured accuracy with Matplotlib plots
- Used OpenCV for image processing steps such as resizing and filtering
Plant Disease Classifier
- Built a CNN in PyTorch using deep learning to classify leaf photos
- Prepared image arrays with NumPy on a Linux machine
""",
    "fullstack": f"""
# Sample Candidate E
{CONTACT}
## Summary
Student who builds complete web applications from the page to the database.
## Education
{EDUCATION}
## Skills
Frontend: JavaScript, HTML, CSS, React
Backend: Node.js, REST API, SQL, MongoDB
Tools: Git, Docker
## Projects
Expense Tracker
- Built the interface with React, HTML and CSS and a Node.js server exposing a REST API
- Stored data in MongoDB and wrote SQL reports for monthly summaries
Event Booking Site
- Wrote JavaScript for the booking flow, kept the code in Git and ran the app in Docker
""",
    "ai_engineer": f"""
# Sample Candidate F
{CONTACT}
## Summary
Student building assistants that answer questions from documents using large language models.
## Education
{EDUCATION}
## Skills
Languages: Python
AI: LLM, RAG, LangChain, prompt engineering, vector databases, deep learning
Backend: FastAPI, REST API, Git
## Projects
Document Question Answering Assistant
- Used retrieval augmented generation (RAG) with LangChain and a vector database to answer questions from PDFs
- Wrote prompt engineering experiments to improve answers and exposed the assistant through a FastAPI REST API
- Managed the code in Git
""",
    "fresher_sparse": f"""
# Sample Candidate G
{CONTACT}
## Summary
First year student who is new to programming and keen to learn.
## Education
B.Tech in Computer Science and Engineering, Sample Institute of Technology, 2025 - 2029
## Skills
Python basics, HTML
## Activities
Member of the college coding club and organised two workshops for juniors
""",
    "tricky_wording": f"""
# Sample Candidate H
{CONTACT}
## Summary
Team player who can excel at learning quickly and react calmly to feedback.
## Education
{EDUCATION}
## Skills
Python, JavaScript (not Java), NoSQL databases, PL/SQL, scikit learn
## Projects
Image Tagger
- Built a machine-learning pipeline and a deep-learning model for tagging photos
- Presented the results as data-visualization dashboards
""",
    "keyword_stuffed": f"""
# Sample Candidate I
{CONTACT}
## Skills
{" ".join(["Python SQL Pandas NumPy Machine Learning scikit-learn Deep Learning TensorFlow PyTorch NLP OpenCV Docker AWS Git React Node.js"] * 6)}
""",
    "ml_engineer_personal": f"""
# Sample Candidate B
{CONTACT}
Gender: Female | Age: 22 | Date of birth: 01 January 2004 | Nationality: Indian
Religion: Hindu | Marital status: Single | Photograph attached
## Summary
Student interested in building machine learning models and deploying them as simple services.
## Education
{EDUCATION}
## Skills
Languages: Python
Machine learning: scikit-learn, Pandas, NumPy, MLflow
Deployment: FastAPI, Docker, Git, AWS
## Projects
House Price Predictor
- Trained regression models with scikit-learn and tracked the experiments with MLflow
- Wrapped the best model in a FastAPI service and packaged it with Docker
Customer Churn Model
- Compared several machine learning classifiers using cross validation and Pandas for feature engineering
- Deployed the service on AWS and kept the code in Git with a CI/CD pipeline using GitHub Actions
""",
}


def parse_markup(text):
    """Turn the simple markup into a list of (kind, text) pairs."""
    lines = []
    for raw in text.strip().splitlines():
        raw = raw.rstrip()
        if raw.startswith("## "):
            lines.append(("heading", raw[3:]))
        elif raw.startswith("# "):
            lines.append(("title", raw[2:]))
        elif raw.startswith("- "):
            lines.append(("bullet", raw[2:]))
        elif raw:
            lines.append(("text", raw))
    return lines


def build_docx(text):
    """Return a DOCX file (as bytes) for one profile."""
    document = Document()
    for kind, content in parse_markup(text):
        if kind == "title":
            document.add_heading(content, level=0)
        elif kind == "heading":
            document.add_heading(content, level=2)
        elif kind == "bullet":
            document.add_paragraph(content, style="List Bullet")
        else:
            document.add_paragraph(content)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(text):
    """Return a PDF file (as bytes) for one profile."""
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=14)
    styles = {
        "title": ParagraphStyle("title", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22),
        "heading": ParagraphStyle("heading", parent=base, fontName="Helvetica-Bold", fontSize=12, spaceBefore=8),
        "bullet": ParagraphStyle("bullet", parent=base, leftIndent=14, bulletIndent=2),
        "text": base,
    }
    story = []
    for kind, content in parse_markup(text):
        bullet = "\u2022" if kind == "bullet" else None
        story.append(Paragraph(escape(content), styles[kind], bulletText=bullet))
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm).build(story)
    return buffer.getvalue()


# ---------- awkward layouts and broken files ----------


def build_docx_nested_table():
    """A DOCX whose skills sit in a table inside a table cell (a common template layout)."""
    document = Document()
    document.add_heading("Sample Candidate J", level=0)
    document.add_paragraph(f"{CONTACT}")
    document.add_paragraph(
        "Student with two data projects and an interest in machine learning engineering roles."
    )
    outer = document.add_table(rows=1, cols=2)
    left, right = outer.rows[0].cells
    left.paragraphs[0].text = "Skills"
    inner = left.add_table(rows=2, cols=2)
    for cell, skill in zip(
        [c for row in inner.rows for c in row.cells], ["Python", "Docker", "Git", "FastAPI"]
    ):
        cell.text = skill
    right.paragraphs[0].text = "Projects"
    right.add_paragraph("Built a machine learning service and packaged it for deployment.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_docx_textbox():
    """A DOCX whose skills sidebar is a text box, as in many designed templates."""
    document = Document()
    document.add_heading("Sample Candidate K", level=0)
    document.add_paragraph(f"{CONTACT}")
    document.add_paragraph(
        "Final year student looking for entry level data roles, with two analysis projects "
        "completed during the degree."
    )
    textbox = parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox>'
        "<w:txbxContent>"
        "<w:p><w:r><w:t>Skills</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Python, SQL, Pandas, Power BI, Excel</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    document.element.body.insert(len(document.element.body) - 1, textbox)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_scanned_pdf():
    """A PDF that contains only a picture of text, like a scanned resume."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (600, 200), "white")
    ImageDraw.Draw(image).text((20, 80), "Python SQL Pandas machine learning", fill="black")
    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")
    image_buffer.seek(0)

    from reportlab.lib.utils import ImageReader

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    page.drawImage(ImageReader(image_buffer), 50, 600, width=480, height=160)
    page.save()
    return buffer.getvalue()


def main():
    OUTPUT_FOLDER.mkdir(exist_ok=True)
    files = {}
    for name, text in PROFILES.items():
        files[f"{name}.docx"] = build_docx(text)
        files[f"{name}.pdf"] = build_pdf(text)
    files["layout_nested_table.docx"] = build_docx_nested_table()
    files["layout_textbox.docx"] = build_docx_textbox()
    files["broken_scanned.pdf"] = build_scanned_pdf()
    for filename, data in files.items():
        (OUTPUT_FOLDER / filename).write_bytes(data)
    print(f"Wrote {len(files)} files to {OUTPUT_FOLDER}")


if __name__ == "__main__":
    main()