"""Resume parsing: validate an uploaded file and extract its text.

Files are processed in memory only. Nothing is written to disk.
"""

import io
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import DependencyError

MAX_FILE_SIZE_MB = 5
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MIN_TEXT_LENGTH = 50  # fewer characters than this probably means no real text


class ResumeParseError(Exception):
    """Raised with a student-friendly message when a resume cannot be read."""


def validate_file(filename, size_bytes):
    """Check the file type and size. Return the lowercase extension if valid."""
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            "Unsupported file type. Please upload a PDF or DOCX resume."
        )
    if size_bytes == 0:
        raise ResumeParseError("The file is empty. Please upload a different file.")
    if size_bytes > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise ResumeParseError(
            f"The file is larger than {MAX_FILE_SIZE_MB} MB. "
            "Please upload a smaller file."
        )
    return extension


def extract_text_from_pdf(file_bytes):
    """Return the text from every page of a PDF."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise ResumeParseError(
                "This PDF is password protected. Please remove the password "
                "and upload it again."
            )
        pages = [page.extract_text() or "" for page in reader.pages]
    except ResumeParseError:
        raise
    except DependencyError as error:
        raise ResumeParseError(
            "This PDF uses an encryption method that this app cannot read. Please "
            "save or print a copy without a password or restrictions and upload that copy."
        ) from error
    except Exception as error:
        raise ResumeParseError(
            "This PDF could not be read. It may be corrupted."
        ) from error
    return "\n".join(pages)


def _container_texts(container, with_text_boxes=True):
    """Return the text of a body, table cell, header or footer.

    Reads paragraphs, tables (including tables inside table cells) and, for the
    outermost container, text boxes, which python-docx does not read by itself.
    """
    texts = [paragraph.text for paragraph in container.paragraphs]
    for table in container.tables:  # many resumes use tables for layout
        for row in table.rows:
            for cell in row.cells:
                texts.extend(_container_texts(cell, with_text_boxes=False))
    if with_text_boxes:  # designed templates often put skills in text boxes
        for xml_paragraph in container._element.xpath(".//w:txbxContent//w:p"):
            texts.append(Paragraph(xml_paragraph, container).text)
    return texts


def extract_text_from_docx(file_bytes):
    """Return the text of a DOCX file: body, nested tables, text boxes, headers, footers."""
    try:
        document = Document(io.BytesIO(file_bytes))
        texts = _container_texts(document._body)
        for section in document.sections:
            texts.extend(_container_texts(section.header))
            texts.extend(_container_texts(section.footer))
    except Exception as error:
        raise ResumeParseError(
            "This DOCX file could not be read. It may be corrupted."
        ) from error
    return "\n".join(texts)


def extract_resume_text(filename, file_bytes):
    """Validate the file, extract its text and return it.

    Raises ResumeParseError with a readable message if anything goes wrong.
    """
    extension = validate_file(filename, len(file_bytes))

    if extension == ".pdf":
        text = extract_text_from_pdf(file_bytes)
    else:
        text = extract_text_from_docx(file_bytes)

    text = text.strip()
    if len(text) < MIN_TEXT_LENGTH:
        raise ResumeParseError(
            "We could not find enough text in this file. If it is a scanned "
            "image, please upload a text-based PDF or a DOCX instead."
        )
    return text